"""Convert 本科专业目录.docx to Markdown."""
import docx
import re
import io

INPUT = r'D:\vscode.project\.claude\skills\咨询侠\知识库\数据层\本科专业目录.docx'
OUTPUT = r'D:\vscode.project\.claude\skills\咨询侠\知识库\数据层\本科专业目录.md'

doc = docx.Document(INPUT)

lines = []
for para in doc.paragraphs:
    text = para.text.strip()
    if not text:
        lines.append('')
        continue

    style = para.style.name if para.style else ''

    # Detect heading levels
    if 'Heading' in style or 'heading' in style:
        m = re.search(r'\d+', style)
        level = min(int(m.group()) + 1, 4) if m else 2  # shift down to reserve # for doc title
        lines.append('#' * level + ' ' + text)
    else:
        lines.append(text)

# Extract tables
for i, table in enumerate(doc.tables):
    lines.append('')
    lines.append('')
    for row_idx, row in enumerate(table.rows):
        cells = [cell.text.strip().replace('\n', '<br>') for cell in row.cells]
        lines.append('| ' + ' | '.join(cells) + ' |')
        if row_idx == 0:
            lines.append('| ' + ' | '.join(['---'] * len(cells)) + ' |')

result = '\n'.join(lines)

# Write with explicit UTF-8 encoding (no BOM)
with io.open(OUTPUT, 'w', encoding='utf-8') as f:
    f.write(result)

print(f'Done. {len(doc.paragraphs)} paragraphs, {len(doc.tables)} tables, {len(result)} chars.')