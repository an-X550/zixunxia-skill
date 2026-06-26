"""
咨询侠 — AI复盘报告 DOCX 导出工具
用法: python export_review_docx.py '<json_input>' 或传入 JSON 文件路径
"""
import json
import sys
import os
from datetime import date
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn


def set_cell_shading(cell, color_hex):
    shading = cell._element.get_or_add_tcPr()
    shading_elm = shading.makeelement(qn('w:shd'), {
        qn('w:fill'): color_hex,
        qn('w:val'): 'clear'
    })
    shading.append(shading_elm)


def add_para(doc, text, bold=False, size=10.5, color=None, alignment=None, space_after=6, italic=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = '微软雅黑'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    run.font.size = Pt(size)
    if bold:
        run.bold = True
    if italic:
        run.italic = True
    if color:
        run.font.color.rgb = RGBColor(*color)
    if alignment is not None:
        p.alignment = alignment
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    return p


def make_table(doc, headers, rows, header_color='4472C4', first_row_highlight=False):
    """创建格式化表格"""
    ncols = len(headers)
    table = doc.add_table(rows=len(rows) + 1, cols=ncols, style='Light Grid Accent 1')
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # 表头
    for j, h in enumerate(headers):
        cell = table.rows[0].cells[j]
        cell.paragraphs[0].clear()
        run = cell.paragraphs[0].add_run(h)
        run.font.name = '微软雅黑'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
        run.bold = True
        run.font.size = Pt(9)
        set_cell_shading(cell, header_color)
        run.font.color.rgb = RGBColor(255, 255, 255)

    # 数据行
    for i, row_data in enumerate(rows):
        for j, val in enumerate(row_data):
            cell = table.rows[i + 1].cells[j]
            cell.paragraphs[0].clear()
            run = cell.paragraphs[0].add_run(str(val))
            run.font.name = '微软雅黑'
            run._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
            run.font.size = Pt(9)
            if first_row_highlight and i == 0:
                set_cell_shading(cell, 'E8F0FE')

    return table


def generate_review_docx(data):
    doc = Document()

    style = doc.styles['Normal']
    font = style.font
    font.name = '微软雅黑'
    font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    student_name = data.get('student_name', '考生')
    doc_date = data.get('date', date.today().isoformat())
    major_category = data.get('major_category', '')
    form = data.get('form', '文字咨询')
    duration = data.get('duration', '')
    province = data.get('province', '')
    score = data.get('score', '')
    rank = data.get('rank', '')
    subject = data.get('subject', '')

    # ===== 封面信息 =====
    title = doc.add_heading(f'咨询侠复盘报告', level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run(f'{student_name} — {major_category}（{form}）')
    run.font.name = '微软雅黑'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(80, 80, 80)

    # 基本信息表
    doc.add_paragraph()
    info_headers = ['项目', '内容']
    info_rows = [
        ['考生昵称', student_name],
        ['咨询日期', doc_date],
        ['咨询形式', form],
        ['专业大类', major_category],
        ['省份/科类', f'{province} · {subject}'],
        ['成绩/位次', f'{score}分 / 省排{rank}'],
        ['咨询时长', duration],
    ]
    make_table(doc, info_headers, info_rows, header_color='666666')
    doc.add_paragraph()

    # ===== 第一部分 =====
    doc.add_heading('一、考生的问题是什么？根本矛盾点是什么？', level=1)

    part1 = data.get('part1', {})
    guiding_q1 = part1.get('guiding_questions', [])
    professional_supplement1 = part1.get('professional_supplement', '')

    doc.add_heading('【引导提问】', level=2)
    if isinstance(guiding_q1, list):
        for q in guiding_q1:
            p = doc.add_paragraph(style='List Bullet')
            run = p.add_run(q)
            run.font.size = Pt(10.5)
    else:
        add_para(doc, str(guiding_q1), size=10.5)

    doc.add_heading('【专业补充】', level=2)
    add_para(doc, professional_supplement1, size=10.5)

    # ===== 第二部分 =====
    doc.add_heading('二、我是如何解决的？', level=1)

    part2 = data.get('part2', {})
    guiding_q2 = part2.get('guiding_questions', [])
    means_end_table = part2.get('means_end_table', {})

    doc.add_heading('【引导提问】', level=2)
    if isinstance(guiding_q2, list):
        for q in guiding_q2:
            p = doc.add_paragraph(style='List Bullet')
            run = p.add_run(q)
            run.font.size = Pt(10.5)
    else:
        add_para(doc, str(guiding_q2), size=10.5)

    doc.add_heading('【专业补充】', level=2)
    add_para(doc, '以下是从聊天记录中梳理的"手段-目的"链条：', size=10, color=(100, 100, 100))
    if means_end_table:
        make_table(doc,
            means_end_table.get('headers', ['目的', '手段', '效果']),
            means_end_table.get('rows', []),
            header_color='4472C4')

    # ===== 第三部分 =====
    doc.add_heading('三、解决的怎么样？哪里做得好，哪里做得不好？', level=1)

    part3 = data.get('part3', {})
    guiding_q3 = part3.get('guiding_questions', [])
    good_table = part3.get('good_table', {})
    bad_table = part3.get('bad_table', {})

    doc.add_heading('【引导提问】', level=2)
    if isinstance(guiding_q3, list):
        for q in guiding_q3:
            p = doc.add_paragraph(style='List Bullet')
            run = p.add_run(q)
            run.font.size = Pt(10.5)

    doc.add_heading('【专业补充 — 做得好的】', level=2)
    if good_table:
        make_table(doc,
            good_table.get('headers', ['目的', '手段', '为什么好']),
            good_table.get('rows', []),
            header_color='2E7D32')

    doc.add_heading('【专业补充 — 可以做得更好的】', level=2)
    if bad_table:
        make_table(doc,
            bad_table.get('headers', ['目的', '手段', '问题在哪', '原因分析']),
            bad_table.get('rows', []),
            header_color='C62828')

    # ===== 第四部分 =====
    doc.add_heading('四、下次遇到同样或类似的问题，我会怎么做？', level=1)

    part4 = data.get('part4', {})
    guiding_q4 = part4.get('guiding_questions', [])
    suggestions = part4.get('suggestions', [])
    voice_optimization = part4.get('voice_optimization', {})

    doc.add_heading('【引导提问】', level=2)
    if isinstance(guiding_q4, list):
        for q in guiding_q4:
            p = doc.add_paragraph(style='List Bullet')
            run = p.add_run(q)
            run.font.size = Pt(10.5)

    doc.add_heading('【专业补充】', level=2)
    if suggestions:
        for i, s in enumerate(suggestions, 1):
            add_para(doc, f'{i}. {s}', size=10.5, bold=False)

    if voice_optimization:
        doc.add_heading('语音咨询额外建议', level=3)
        make_table(doc,
            voice_optimization.get('headers', ['阶段', '优化建议']),
            voice_optimization.get('rows', []),
            header_color='6A1B9A')

    # ===== 页脚 =====
    doc.add_paragraph()
    add_para(doc, '—— 问问大象 · 咨询侠 ——', size=8, color=(150, 150, 150), alignment=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, '引导提问 + 专业补充 · 先让咨询侠自己想，再给补充', size=8, color=(150, 150, 150), alignment=WD_ALIGN_PARAGRAPH.CENTER)

    return doc


def main():
    if len(sys.argv) > 1:
        raw = sys.argv[1]
        try:
            data = json.loads(raw)
        except json.JSONDecodeError:
            filepath = sys.argv[1]
            if os.path.exists(filepath):
                with open(filepath, 'r', encoding='utf-8') as f:
                    data = json.load(f)
            else:
                print(f"Error: invalid JSON or file not found: {filepath}")
                sys.exit(1)
    else:
        raw = sys.stdin.read()
        data = json.loads(raw)

    doc = generate_review_docx(data)

    student_name = data.get('student_name', '考生')
    output_dir = os.path.dirname(os.path.abspath(__file__))
    cases_dir = os.path.join(output_dir, 'cases')
    output_path = os.path.join(cases_dir, f'{student_name}-AI复盘报告.docx')

    doc.save(output_path)
    print(f'DOCX saved: {output_path}')


if __name__ == '__main__':
    main()
