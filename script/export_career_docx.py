"""快速将 MD 转为 DOCX — 通用导出"""
import sys, os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

def md_to_docx(md_path, docx_path):
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = '微软雅黑'
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    for s in doc.sections:
        s.top_margin = Cm(2); s.bottom_margin = Cm(2)
        s.left_margin = Cm(2.5); s.right_margin = Cm(2.5)

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_table = False
    for line in lines:
        line = line.rstrip('\n')
        if not line.strip():
            doc.add_paragraph()
            continue

        stripped = line.lstrip()

        # 表格
        if stripped.startswith('|') and stripped.endswith('|'):
            cells = [c.strip() for c in stripped.strip('|').split('|')]
            if all(c.startswith(':--') or c.startswith('--') or c == '' for c in cells if c):
                continue  # 分隔行跳过
            if not in_table:
                in_table = True
                tbl = doc.add_table(rows=0, cols=len(cells), style='Light Grid Accent 1')
            row = tbl.add_row()
            for i, c in enumerate(cells):
                p = row.cells[i].paragraphs[0]
                p.clear()
                run = p.add_run(c)
                run.font.size = Pt(9)
                if '**' in c:
                    c2 = c.replace('**','')
                    p.clear()
                    run = p.add_run(c2)
                    run.bold = True; run.font.size = Pt(9)
        else:
            in_table = False
            # 标题
            if stripped.startswith('# ') and not stripped.startswith('## '):
                h = doc.add_heading(stripped[2:], level=1)
            elif stripped.startswith('## '):
                h = doc.add_heading(stripped[3:], level=2)
            elif stripped.startswith('### '):
                h = doc.add_heading(stripped[4:], level=3)
            elif stripped.startswith('- ') or stripped.startswith('* '):
                p = doc.add_paragraph(stripped[2:], style='List Bullet')
            elif stripped.startswith('> '):
                p = doc.add_paragraph()
                run = p.add_run(stripped[2:])
                run.italic = True; run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(100,100,100)
            elif stripped == '---':
                doc.add_paragraph('—' * 30)
            else:
                p = doc.add_paragraph(stripped)
                # 处理粗体
                for run in p.runs:
                    if '**' in run.text:
                        pass  # 简单处理，略过复杂内联

    doc.save(docx_path)
    print(f'OK: {docx_path}')

if __name__ == '__main__':
    md_path = sys.argv[1]
    docx_path = sys.argv[2] if len(sys.argv) > 2 else md_path.replace('.md', '.docx')
    md_to_docx(md_path, docx_path)