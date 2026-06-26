"""
咨询侠 — 专业推荐文档 DOCX 导出工具
用法: python export_docx.py '<json_input>' 或从 stdin 读取 JSON
"""
import json
import sys
import os
from datetime import date
from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn


def set_cell_shading(cell, color_hex):
    """设置单元格背景色"""
    shading = cell._element.get_or_add_tcPr()
    shading_elm = shading.makeelement(qn('w:shd'), {
        qn('w:fill'): color_hex,
        qn('w:val'): 'clear'
    })
    shading.append(shading_elm)


def add_styled_paragraph(doc, text, style='Normal', bold=False, size=None, color=None, alignment=None, space_after=None):
    """添加格式化段落"""
    p = doc.add_paragraph(style=style)
    run = p.add_run(text)
    if bold:
        run.bold = True
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor(*color)
    if alignment is not None:
        p.alignment = alignment
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    return p


def add_msyh_run(paragraph, text, bold=False, size=None, color=None):
    """在段落中添加微软雅黑格式的 run，确保中文字体正确"""
    run = paragraph.add_run(text)
    run.font.name = '微软雅黑'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    if bold:
        run.bold = True
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor(*color) if isinstance(color, tuple) else color
    return run


def generate_docx(data):
    """基于 JSON 数据生成专业推荐 DOCX"""
    doc = Document()

    # 设置默认字体
    style = doc.styles['Normal']
    font = style.font
    font.name = '微软雅黑'
    font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

    # 窄页边距
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    student_name = data.get('student_name', '考生')
    doc_date = data.get('date', date.today().isoformat())
    majors = data.get('majors', [])
    comparison_table = data.get('comparison_table', {})
    pitfalls = data.get('pitfalls', [])
    summary = data.get('summary', '')

    # ===== 标题 =====
    title = doc.add_heading(f'高考志愿填报专业推荐文档', level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title.runs:
        run.font.name = '微软雅黑'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

    # 副标题
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = add_msyh_run(sub, f'—— 为 {student_name} 定制 ｜ {doc_date}', size=10, color=(100, 100, 100))

    doc.add_paragraph()  # 空行

    # ===== 推荐逻辑摘要 =====
    if summary:
        h = doc.add_heading('推荐逻辑摘要', level=2)
        for run in h.runs:
            run.font.name = '微软雅黑'
            run._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
        add_styled_paragraph(doc, summary, size=10.5)

    # ===== 推荐专业 =====
    h2 = doc.add_heading(f'推荐专业（共 {len(majors)} 个）', level=2)
    for run in h2.runs:
        run.font.name = '微软雅黑'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

    strength_labels = {
        5: '⭐⭐⭐⭐⭐ 最高推荐',
        4: '⭐⭐⭐⭐ 强烈推荐',
        3: '⭐⭐⭐ 推荐',
        2: '⭐⭐ 可考虑',
        1: '⭐ 备选'
    }
    strength_order = ['🥇', '🥈', '🥉', '④', '⑤']

    for i, m in enumerate(majors):
        name = m.get('name', '')
        strength = m.get('strength', 3)
        medal = strength_order[i] if i < len(strength_order) else '•'

        # 专业标题
        h3 = doc.add_heading(f'{medal} {name}', level=3)
        for run in h3.runs:
            run.font.name = '微软雅黑'
            run._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

        # 推荐强度
        strength_text = strength_labels.get(strength, f'推荐度 {strength}/5')
        p = doc.add_paragraph()
        add_msyh_run(p, f'推荐强度：{strength_text}', bold=True, size=10.5)

        # 信息表格
        table = doc.add_table(rows=7, cols=2, style='Light Grid Accent 1')
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        # 设置列宽
        for row in table.rows:
            row.cells[0].width = Cm(3)
            row.cells[1].width = Cm(12)

        fields = [
            ('推荐原因', m.get('reason', '')),
            ('适合点', m.get('fit_points', '')),
            ('需要了解的现实面', m.get('reality_check', '')),
            ('可能发展方向', m.get('development', '')),
            ('赚钱路径类型', m.get('earning_type', '')),
            ('学校层次敏感度', m.get('prestige_sensitivity', '')),
            ('给你的行动建议', m.get('action_advice', '')),
        ]

        # 表头行样式
        for row in table.rows:
            set_cell_shading(row.cells[0], 'E8F0FE')
            row.cells[0].paragraphs[0].clear()
            run = add_msyh_run(row.cells[0].paragraphs[0], '', size=9.5)

        for j, (label, value) in enumerate(fields):
            # 标签
            cell_label = table.rows[j].cells[0]
            cell_label.paragraphs[0].clear()
            add_msyh_run(cell_label.paragraphs[0], label, bold=True, size=9.5)

            # 值
            cell_value = table.rows[j].cells[1]
            cell_value.paragraphs[0].clear()
            add_msyh_run(cell_value.paragraphs[0], str(value), size=9.5)

        doc.add_paragraph()  # 专业间间距

    # ===== 专业对比一览 =====
    if comparison_table:
        h_cmp = doc.add_heading('专业对比一览', level=2)
        for run in h_cmp.runs:
            run.font.name = '微软雅黑'
            run._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
        headers = comparison_table.get('headers', [])
        rows = comparison_table.get('rows', [])

        if headers and rows:
            ncols = len(headers)
            table = doc.add_table(rows=len(rows) + 1, cols=ncols, style='Light Grid Accent 1')
            table.alignment = WD_TABLE_ALIGNMENT.CENTER

            # 表头
            for j, h in enumerate(headers):
                cell = table.rows[0].cells[j]
                cell.paragraphs[0].clear()
                run = add_msyh_run(cell.paragraphs[0], h, bold=True, size=9)
                set_cell_shading(cell, '4472C4')
                run.font.color.rgb = RGBColor(255, 255, 255)

            # 数据行
            for i, row_data in enumerate(rows):
                for j, cell_val in enumerate(row_data):
                    cell = table.rows[i + 1].cells[j]
                    cell.paragraphs[0].clear()
                    add_msyh_run(cell.paragraphs[0], str(cell_val), size=9)
                    if i == 0:
                        set_cell_shading(cell, 'F0F8E8')

    # ===== 避坑清单 =====
    if pitfalls:
        h_pit = doc.add_heading('选专业避坑清单', level=2)
        for run in h_pit.runs:
            run.font.name = '微软雅黑'
            run._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
        p_intro = doc.add_paragraph()
        add_msyh_run(p_intro, '以下提醒基于你的个人画像定制，不是通用套话：', size=9.5, color=(100, 100, 100))

        for idx, pitfall in enumerate(pitfalls, 1):
            p = doc.add_paragraph()
            add_msyh_run(p, f'⚠️ {idx}. {pitfall}', size=10)
            p.paragraph_format.space_after = Pt(6)

    # ===== 页脚 =====
    doc.add_paragraph()
    footer_p = doc.add_paragraph()
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_msyh_run(footer_p, '—— 问问大象 · 咨询侠 ——', size=8, color=(150, 150, 150))
    footer_p2 = doc.add_paragraph()
    footer_p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_msyh_run(footer_p2, '授人以渔，而非授人以鱼', size=8, color=(150, 150, 150))

    return doc


def main():
    if len(sys.argv) > 1:
        raw = sys.argv[1]
        try:
            data = json.loads(raw)
        except json.JSONDecodeError:
            # 可能是文件路径
            filepath = sys.argv[1]
            if os.path.exists(filepath):
                with open(filepath, 'r', encoding='utf-8') as f:
                    data = json.load(f)
            else:
                print(f"Error: invalid JSON or file not found: {filepath}")
                sys.exit(1)
    else:
        # 从 stdin 读取
        raw = sys.stdin.read()
        data = json.loads(raw)

    doc = generate_docx(data)

    # 输出路径
    student_name = data.get('student_name', '考生')
    version = data.get('version', '')
    output_dir = os.path.dirname(os.path.abspath(__file__))
    skill_root = os.path.dirname(output_dir)
    cases_dir = os.path.join(skill_root, 'cases')
    os.makedirs(cases_dir, exist_ok=True)
    ver_suffix = f'-{version}' if version else ''
    output_path = os.path.join(cases_dir, f'{student_name}-专业推荐文档{ver_suffix}.docx')

    doc.save(output_path)
    print(f'DOCX saved: {output_path}')


if __name__ == '__main__':
    main()
